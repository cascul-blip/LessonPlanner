"""One-time import of the old LessonPlan.xlsx workbook.

Each "Week N" sheet has one or more blocks: a header row of class names with
Monday..Friday in column A underneath. Each multi-line cell is split into
lesson / special / homework lines.
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from . import schoolcal
from .db import Database, Entry

SPECIAL_RE = re.compile(
    r"^(pop quiz|quiz\b|test(\s*\d+)?$|test\s*\(|diagnostic test|lab\s*\d|activity\b)",
    re.IGNORECASE,
)
HOMEWORK_RE = re.compile(
    r"^(read|answer|complete|study|bring|begin|preview|report sheet|problems)", re.IGNORECASE
)
HW_MARKER_RE = re.compile(r"^hw\b:?\s*", re.IGNORECASE)
WEEK_SHEET_RE = re.compile(r"^week\s*(\d+)$", re.IGNORECASE)


def split_cell(text: str) -> tuple[str, str, str]:
    """Split an old spreadsheet cell into (lesson, special, homework)."""
    lesson, special, homework = [], [], []
    in_homework = False
    for raw in text.replace("\r", "").split("\n"):
        line = re.sub(r"\bLesosn\b", "Lesson", raw.strip())
        if not line:
            continue
        marker = HW_MARKER_RE.match(line)
        if marker:
            in_homework = True
            line = line[marker.end():]
            if not line:
                continue
        if SPECIAL_RE.match(line):
            special.append(line)
        elif in_homework or HOMEWORK_RE.match(line):
            in_homework = True
            homework.append(line)
        else:
            lesson.append(line)
    return "\n".join(lesson), "\n".join(special), "\n".join(homework)


@dataclass
class ImportReport:
    imported: int = 0
    skipped_existing: int = 0
    no_school: list[date] = field(default_factory=list)
    unknown_headers: set[str] = field(default_factory=set)
    specials: list[tuple[date, str, str]] = field(default_factory=list)

    def summary(self) -> str:
        lines = [f"Imported {self.imported} lessons."]
        if self.skipped_existing:
            lines.append(f"Skipped {self.skipped_existing} cells that already had data.")
        for d in self.no_school:
            lines.append(f"Marked {d:%a %b %d} as No School (no lessons that day).")
        if self.unknown_headers:
            lines.append("Ignored unknown columns: " + ", ".join(sorted(self.unknown_headers)))
        if self.specials:
            lines.append("Special items found:")
            lines += [f"  {d:%a %b %d}  {c}: {s.replace(chr(10), ' / ')}"
                      for d, c, s in self.specials]
        return "\n".join(lines)


def _is_labor_day(day: date) -> bool:
    return day.month == 9 and day.weekday() == 0 and day.day <= 7


def import_workbook(db: Database, path: Path, overwrite: bool = False) -> ImportReport:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    report = ImportReport()
    week1 = db.week1_monday()
    existing_no_school = db.no_school_days()

    with db.conn:
        for ws in wb.worksheets:
            m = WEEK_SHEET_RE.match(ws.title.strip())
            if not m:
                continue
            days = schoolcal.week_days(schoolcal.week_monday(int(m.group(1)), week1))
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            filled_days: set[date] = set()

            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    if c == 0 or not isinstance(value, str) or not value.strip():
                        continue
                    below = [rows[r + i][0] if r + i < len(rows) else None for i in range(1, 6)]
                    if [str(b).strip() if b else "" for b in below] != schoolcal.DAY_NAMES:
                        continue  # not a header cell
                    cls = db.class_by_name(value.strip())
                    if cls is None:
                        report.unknown_headers.add(value.strip())
                        continue
                    for i, day in enumerate(days):
                        cell_row = rows[r + 1 + i]
                        text = cell_row[c] if c < len(cell_row) else None
                        if not isinstance(text, str) or not text.strip():
                            continue
                        filled_days.add(day)
                        lesson, special, homework = split_cell(text)
                        if not overwrite and not db.get_entry(cls.id, day).is_empty():
                            report.skipped_existing += 1
                            continue
                        db._save_entry(Entry(cls.id, day, lesson, special, homework))
                        report.imported += 1
                        if special:
                            report.specials.append((day, cls.name, special))

            if filled_days:
                for day in days:
                    if day not in filled_days and day not in existing_no_school:
                        reason = "Labor Day" if _is_labor_day(day) else "No school"
                        db.conn.execute(
                            "INSERT OR IGNORE INTO no_school(date, reason) VALUES (?, ?)",
                            (day.isoformat(), reason),
                        )
                        report.no_school.append(day)
    wb.close()
    report.specials.sort()
    report.no_school.sort()
    return report


def main(argv: list[str] | None = None) -> int:
    from . import config

    parser = argparse.ArgumentParser(description="Import the old LessonPlan.xlsx workbook.")
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("--db", type=Path, help="database file (default: configured folder)")
    parser.add_argument("--overwrite", action="store_true",
                        help="replace cells that already have data")
    args = parser.parse_args(argv)

    db_file = args.db
    if db_file is None:
        folder = config.db_folder()
        if folder is None:
            print("No database folder configured; pass --db or run the app first.",
                  file=sys.stderr)
            return 1
        db_file = config.db_path(folder)
    db = Database(db_file)
    try:
        print(import_workbook(db, args.xlsx, args.overwrite).summary())
    finally:
        db.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
